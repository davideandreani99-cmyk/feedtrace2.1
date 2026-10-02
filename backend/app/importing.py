"""Import iniziale CSV/XLSX delle anagrafiche con anteprima e conferma (RF-01).

Colonne = nomi dei campi degli schemi (es. `code;name;type;unit`), numeri in formato
italiano (virgola decimale), sì/no, date gg/mm/aaaa. I riferimenti si danno per codice
(es. `site_code`). Un import ha esito solo se TUTTE le righe sono valide (transazione unica).
"""

import csv
import io
import types
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from enum import StrEnum
from typing import Any, Union, get_args, get_origin

from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError
from sqlalchemy import select

from app import audit
from app.deps import Ctx
from app.models import Container, DestinationLot, RawMaterial, Recipe, Site, Supplier
from app.routers.recipes import create_recipe_version
from app.schemas import (
    ContainerIn,
    DestinationLotIn,
    RawMaterialIn,
    RecipeIn,
    RecipeLineIn,
    SupplierIn,
)
from domain.parsing import parse_bool_it, parse_date_it, parse_decimal_it

MAX_FILE_BYTES = 5 * 1024 * 1024


@dataclass
class RowError:
    line: int
    message: str


@dataclass
class ImportResult:
    entity: str
    total: int = 0
    ok: int = 0
    errors: list[RowError] = field(default_factory=list)
    applied: bool = False


# --- lettura file ------------------------------------------------------------------
def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def read_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header = [_cell(h).lower() for h in next(iterator, ())]
        rows = []
        for raw in iterator:
            row = {header[i]: _cell(v) for i, v in enumerate(raw) if i < len(header) and header[i]}
            if any(row.values()):
                rows.append(row)
        return rows
    if name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        first_line = text.splitlines()[0] if text.strip() else ""
        delimiter = ";" if first_line.count(";") >= first_line.count(",") else ","
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        rows = []
        for raw in reader:
            row = {(k or "").strip().lower(): (v or "").strip() for k, v in raw.items() if k}
            if any(row.values()):
                rows.append(row)
        return rows
    raise ValueError("Formato non supportato: usa un file .csv o .xlsx")


# --- conversione valori ------------------------------------------------------------
def _base_type(annotation: Any) -> Any:
    if get_origin(annotation) in (Union, types.UnionType):
        args = [a for a in get_args(annotation) if a is not type(None)]
        if len(args) == 1:
            return args[0]
    return annotation


def coerce(schema: type[BaseModel], row: dict[str, str]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for name, info in schema.model_fields.items():
        raw = row.get(name)
        if raw is None or raw == "":
            continue
        base = _base_type(info.annotation)
        try:
            if base is Decimal:
                out[name] = parse_decimal_it(raw)
            elif base is bool:
                out[name] = parse_bool_it(raw)
            elif base is date:
                out[name] = parse_date_it(raw)
            elif isinstance(base, type) and issubclass(base, StrEnum):
                out[name] = _enum_value(base, raw)
            else:
                out[name] = raw
        except ValueError as exc:
            raise ValueError(f"{name}: {exc}") from exc
    return out


def _enum_value(enum_cls: type[StrEnum], raw: str) -> StrEnum:
    for candidate in (raw, raw.upper(), raw.lower()):
        try:
            return enum_cls(candidate)
        except ValueError:
            continue
    allowed = ", ".join(member.value for member in enum_cls)
    raise ValueError(f"valore non ammesso {raw!r} (ammessi: {allowed})")


def _format_validation(exc: ValidationError) -> str:
    return "; ".join(
        f"{'.'.join(str(p) for p in err['loc']) or 'riga'}: {err['msg']}" for err in exc.errors()
    )


# --- specifiche per entità ---------------------------------------------------------
@dataclass(frozen=True)
class Spec:
    schema: type[BaseModel]
    model: type
    key: tuple[str, ...]
    # (colonna nel file, modello riferito, attributo codice, campo di destinazione)
    lookups: tuple[tuple[str, type, str, str], ...] = ()


SPECS: dict[str, Spec] = {
    "suppliers": Spec(SupplierIn, Supplier, ("name",)),
    "raw-materials": Spec(RawMaterialIn, RawMaterial, ("code",)),
    "containers": Spec(
        ContainerIn, Container, ("site_id", "code"), (("site_code", Site, "asl_code", "site_id"),)
    ),
    "destination-lots": Spec(
        DestinationLotIn,
        DestinationLot,
        ("internal_code",),
        (("site_code", Site, "asl_code", "site_id"),),
    ),
}
ENTITIES = [*SPECS.keys(), "recipes"]


def _norm(value: Any) -> str:
    return str(value).strip().lower()


def _code_map(ctx: Ctx, model: type, attr: str) -> dict[str, Any]:
    rows = ctx.db.execute(
        select(model.id, getattr(model, attr)).where(
            model.tenant_id == ctx.tenant_id, model.active.is_(True)
        )
    ).all()
    return {_norm(code): id_ for id_, code in rows}


def run_import(
    ctx: Ctx, entity: str, rows: list[dict[str, str]], *, confirm: bool
) -> ImportResult:
    if entity == "recipes":
        return _import_recipes(ctx, rows, confirm=confirm)
    spec = SPECS[entity]
    result = ImportResult(entity=entity, total=len(rows))
    lookup_maps = {
        col: _code_map(ctx, ref_model, attr) for col, ref_model, attr, _ in spec.lookups
    }
    existing = {
        tuple(_norm(v) for v in row)
        for row in ctx.db.execute(
            select(*[getattr(spec.model, k) for k in spec.key]).where(
                spec.model.tenant_id == ctx.tenant_id
            )
        ).all()
    }
    seen: set[tuple[str, ...]] = set()
    pending: list[BaseModel] = []

    for line, row in enumerate(rows, start=2):  # riga 1 = intestazione
        try:
            data = coerce(spec.schema, row)
            for col, _model, _attr, target in spec.lookups:
                code = row.get(col, "")
                if not code:
                    raise ValueError(f"colonna obbligatoria mancante: {col}")
                ref_id = lookup_maps[col].get(_norm(code))
                if ref_id is None:
                    raise ValueError(f"{col} '{code}' non trovato")
                data[target] = ref_id
            item = spec.schema.model_validate(data)
            key = tuple(_norm(getattr(item, k)) for k in spec.key)
            if key in existing or key in seen:
                raise ValueError("elemento già presente (duplicato)")
            seen.add(key)
            pending.append(item)
        except (ValueError, ValidationError) as exc:
            message = _format_validation(exc) if isinstance(exc, ValidationError) else str(exc)
            result.errors.append(RowError(line, message))

    result.ok = len(pending)
    if confirm and not result.errors and pending:
        for item in pending:
            obj = spec.model(tenant_id=ctx.tenant_id, **audit.plain(item.model_dump()))
            ctx.db.add(obj)
            ctx.db.flush()
            audit.record(
                ctx, entity, obj.id, "IMPORT_CREATE", None, item.model_dump(mode="json")
            )
        ctx.db.commit()
        result.applied = True
    return result


def _import_recipes(ctx: Ctx, rows: list[dict[str, str]], *, confirm: bool) -> ImportResult:
    """Formato lungo: una riga per ingrediente (recipe_code, recipe_name, plant_type, ...,
    material_code, kg_per_t). Crea la versione 1 di ogni ricetta."""
    result = ImportResult(entity="recipes", total=len(rows))
    materials = _code_map(ctx, RawMaterial, "code")
    existing_codes = {
        _norm(c)
        for (c,) in ctx.db.execute(
            select(Recipe.code).where(Recipe.tenant_id == ctx.tenant_id)
        ).all()
    }
    grouped: dict[str, dict[str, Any]] = {}

    for line, row in enumerate(rows, start=2):
        try:
            code = row.get("recipe_code", "")
            if not code:
                raise ValueError("colonna obbligatoria mancante: recipe_code")
            if _norm(code) in existing_codes:
                raise ValueError(f"ricetta '{code}' già presente")
            material_code = row.get("material_code", "")
            material_id = materials.get(_norm(material_code))
            if material_id is None:
                raise ValueError(f"material_code '{material_code}' non trovato")
            line_in = RecipeLineIn.model_validate(
                {
                    "material_id": material_id,
                    "kg_per_t": parse_decimal_it(row.get("kg_per_t", "")),
                }
            )
            entry = grouped.setdefault(
                _norm(code), {"row": row, "code": code, "lines": [], "line": line}
            )
            if any(prev.material_id == material_id for prev in entry["lines"]):
                raise ValueError(f"materia prima '{material_code}' ripetuta nella ricetta")
            entry["lines"].append(line_in)
        except (ValueError, ValidationError) as exc:
            message = _format_validation(exc) if isinstance(exc, ValidationError) else str(exc)
            result.errors.append(RowError(line, message))

    recipes: list[RecipeIn] = []
    for entry in grouped.values():
        row = entry["row"]
        try:
            data = coerce(RecipeIn, row)
            data["code"] = entry["code"]
            data["lines"] = entry["lines"]
            recipes.append(RecipeIn.model_validate(data))
        except (ValueError, ValidationError) as exc:
            message = _format_validation(exc) if isinstance(exc, ValidationError) else str(exc)
            result.errors.append(RowError(entry["line"], message))

    result.ok = len(recipes)
    if confirm and not result.errors and recipes:
        for payload in recipes:
            recipe = create_recipe_version(
                ctx, payload, version=1, valid_from=payload.valid_from or date.today()
            )
            audit.record(
                ctx, "recipes", recipe.id, "IMPORT_CREATE", None, {"code": payload.code}
            )
        ctx.db.commit()
        result.applied = True
    return result
